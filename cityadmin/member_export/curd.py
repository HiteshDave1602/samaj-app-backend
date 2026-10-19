import io
from datetime import datetime, timezone
from uuid import uuid4

import openpyxl
from msoffcrypto.format.ooxml import OOXMLFile
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from cityadmin.member_export.model import ExportJob
from cityadmin.member_export.schema import ExportIn
from core.security import fernet
from member.profile.curd import age
from member.profile.model import Household, Member
from superadmin.audit_log import curd as audit
from superadmin.city_admin.model import AdminUser


def _decrypt(value: bytes | None) -> str | None:
    return fernet.decrypt(value).decode() if value else None


def _rows(db: Session, admin: AdminUser, data: ExportIn) -> list[dict]:
    f = data.filters
    years = func.date_part("year", func.age(Member.dob))
    q = (select(Member, Household.family_name, Household.native_place)
         .outerjoin(Household, Household.household_id == Member.household_id)
         .where(Member.city_id == admin.city_id, Member.deleted_at.is_(None), Member.status.in_(f.status)))
    if f.gender:
        q = q.where(Member.gender == f.gender)
    if f.age_min is not None:
        q = q.where(years >= f.age_min)
    if f.age_max is not None:
        q = q.where(years <= f.age_max)
    if f.family_name:
        q = q.where(Household.family_name.icontains(f.family_name, autoescape=True))
    out = []
    # ponytail: whole city in memory, and `area` matched after decrypting addresses; fine for a city's
    # few thousand members. Needs a searchable area column if cities reach ~100k.
    for m, family_name, native_place in db.execute(q.order_by(Member.full_name, Member.member_id)):
        address = _decrypt(m.address_enc)
        if f.area and f.area.casefold() not in f"{address or ''} {native_place or ''}".casefold():
            continue
        out.append({"full_name": m.full_name, "gender": m.gender, "dob": m.dob, "age": age(m.dob),
                    "mobile": _decrypt(m.mobile_enc), "address": address, "family_name": family_name,
                    "native_place": native_place, "status": m.status, "member_since": m.created_at.date()})
    return out


def _xlsx(columns: list[str], rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    ws.append(columns)
    for r in rows:
        ws.append([r[c] for c in columns])
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":  # a name like "=HYPERLINK(...)" must stay text, not become a formula
                cell.data_type = "s"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export(db: Session, admin: AdminUser, actor: audit.Actor, data: ExportIn) -> tuple[bytes, int]:
    """FR-82/83: the filtered members as an XLSX encrypted with the admin's password (AES-256, opens in Excel).
    The password exists only for this call; neither it nor the file is kept (FR-85)."""
    columns = list(dict.fromkeys(data.columns))
    rows = _rows(db, admin, data)
    encrypted = io.BytesIO()
    OOXMLFile(io.BytesIO(_xlsx(columns, rows))).encrypt(data.password.get_secret_value(), encrypted)
    job = ExportJob(export_job_id=uuid4(), city_id=admin.city_id, exported_by_admin_id=admin.admin_id,
                    filters=data.filters.model_dump(mode="json"), columns_exported=columns, row_count=len(rows),
                    file_format="XLSX")
    db.add(job)
    audit.log(db, actor, "MEMBER_EXPORT", "export_job", job.export_job_id, admin.city_id, None,
              {"filters": job.filters, "columns": columns, "row_count": len(rows)})
    db.commit()
    return encrypted.getvalue(), len(rows)


def list_jobs(db: Session, admin: AdminUser, limit: int, offset: int) -> dict:
    q = select(ExportJob).where(ExportJob.city_id == admin.city_id)
    total = db.scalar(select(func.count()).select_from(q.subquery()))
    jobs = db.scalars(q.order_by(ExportJob.created_at.desc()).limit(limit).offset(offset))
    return {"total": total, "items": [{c.name: getattr(j, c.name) for c in ExportJob.__table__.columns}
                                      for j in jobs]}


def filename(now: datetime | None = None) -> str:
    return f"samaj-members-{(now or datetime.now(timezone.utc)).strftime('%Y%m%d-%H%M')}.xlsx"
