import csv
import io
import json
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from uuid import UUID, uuid4

import openpyxl
from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from cityadmin.member_import.model import ImportJob, ImportRow
from core import private_files
from core.security import fernet
from member.auth.curd import NOTICE_VERSION, mobile_hash
from member.auth.schema import to_e164
from member.profile.model import Consent, Household, Member
from superadmin.audit_log import curd as audit
from superadmin.city_admin.model import AdminUser

# FR-76: the published template. Change HEADERS -> bump TEMPLATE_VERSION.
TEMPLATE_VERSION = "v1"
HEADERS = ["full_name", "mobile", "gender", "dob", "family_name", "native_place", "address"]
EXAMPLE = ["Ramesh Patel", "9876543210", "MALE", "1975-08-15", "Patel", "Surat", "12 Gandhi Road"]
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000  # import_job.total_rows check
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
# ponytail: payload_enc and the uploaded file are kept; the DDD wants them wiped 30 days after commit
# by the retention job (SF-15).


def template(file_format: str) -> bytes:
    if file_format == "CSV":
        buf = io.StringIO()
        csv.writer(buf).writerows([HEADERS, EXAMPLE])
        return buf.getvalue().encode("utf-8-sig")
    wb = openpyxl.Workbook()
    wb.active.append(HEADERS)
    wb.active.append(EXAMPLE)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(v) -> str:
    """An XLSX cell as the text a CSV would have held."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    if isinstance(v, float) and v.is_integer():  # a mobile number typed into a number cell
        return str(int(v))
    return str(v).strip()


def _read(filename: str, data: bytes) -> tuple[str, list[list[str]]]:
    ext = Path(filename or "").suffix.lower()
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(400, "Save the CSV as UTF-8 and upload it again")
        return "CSV", [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    if ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = [[_text(v) for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
            wb.close()
        except Exception:  # openpyxl raises many types for a damaged or mislabelled file
            raise HTTPException(400, "This .xlsx file can't be read. Save it again from Excel and upload it.")
        return "XLSX", rows
    raise HTTPException(400, "Upload a .csv or .xlsx file")


def _date(s: str) -> date | None:
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _clean(cells: list[str]) -> tuple[dict | None, str | None]:
    """One row -> (payload, None) or (None, every problem in it)."""
    v = dict(zip(HEADERS, cells + [""] * len(HEADERS)))
    name, family = " ".join(v["full_name"].split()), " ".join(v["family_name"].split())
    gender, dob = v["gender"].upper(), _date(v["dob"])
    errors = []
    if not 1 <= len(name) <= 120:
        errors.append("full_name is required (max 120 characters)")
    try:
        mobile = to_e164(v["mobile"])
    except ValueError:
        errors.append("mobile is not a valid mobile number")
    if gender not in ("MALE", "FEMALE", "OTHER"):
        errors.append("gender must be MALE, FEMALE or OTHER")
    if dob is None or not date(1900, 1, 1) <= dob <= date.today():
        errors.append("dob must be a past date like 1975-08-15")
    if not 1 <= len(family) <= 120:
        errors.append("family_name is required (max 120 characters)")
    if len(v["native_place"]) > 120 or len(v["address"]) > 500:
        errors.append("native_place is limited to 120 characters and address to 500")
    if any(cells[len(HEADERS):]):
        errors.append("this row has values beyond the template's columns")
    if errors:
        return None, "; ".join(errors)
    return {"full_name": name, "mobile": mobile, "gender": gender, "dob": dob.isoformat(), "family_name": family,
            "native_place": v["native_place"] or None, "address": v["address"] or None}, None


def _same_name(a: str | None, b: str) -> bool:
    return " ".join((a or "").split()).casefold() == b.casefold()


def _job_out(job: ImportJob) -> dict:
    return {c.name: getattr(job, c.name) for c in ImportJob.__table__.columns}


def upload(db: Session, admin: AdminUser, filename: str, data: bytes) -> dict:
    """Validate and deduplicate into a dry-run report (FR-79). Nothing touches members until commit()."""
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "The file is larger than 10 MB")
    file_format, rows = _read(filename, data)
    header = rows[0] if rows else []
    while header and not header[-1]:
        header.pop()
    if header != HEADERS:  # FR-76: whole file refused, nothing written
        raise HTTPException(400, f"Column headers must be exactly: {', '.join(HEADERS)} (template "
                                 f"{TEMPLATE_VERSION}). Download the template. Nothing was imported.")
    body = [(n, r) for n, r in enumerate(rows[1:], start=2) if any(r)]
    if not body:
        raise HTTPException(400, "The file has no member rows")
    if len(body) > MAX_ROWS:
        raise HTTPException(400, f"At most {MAX_ROWS} rows per file; split it and upload the parts")

    cleaned, first_row = [], {}  # first_row: mobile_hash -> row it first appeared in
    for n, cells in body:
        payload, error = _clean(cells)
        mh = mobile_hash(payload["mobile"]) if payload else None
        if mh in first_row:
            payload, error, mh = None, f"Same mobile number as row {first_row[mh]}", None
        elif mh:
            first_row[mh] = n
        cleaned.append((n, payload, error, mh))

    existing, hashes = {}, list(first_row)  # FR-77: match on mobile number, in batches
    for i in range(0, len(hashes), 1000):
        for m in db.execute(select(Member.member_id, Member.mobile_hash, Member.city_id, Member.full_name,
                                   Member.dob, Member.deleted_at).where(Member.mobile_hash.in_(hashes[i:i + 1000]))):
            existing[m.mobile_hash] = m

    job = ImportJob(import_job_id=uuid4(), city_id=admin.city_id, uploaded_by_admin_id=admin.admin_id,
                    file_key=private_files.save(data), file_format=file_format, template_version=TEMPLATE_VERSION,
                    status="DRY_RUN_READY", total_rows=len(body))
    counts, import_rows = Counter(), []
    for n, payload, error, mh in cleaned:
        match = existing.get(mh)
        if payload is None:
            outcome = "ERROR"
        elif match is None:
            outcome = "NEW"
        elif match.city_id != admin.city_id or match.deleted_at is not None:  # FR-71: never another city's member
            outcome, error, payload, match = "ERROR", "This mobile number belongs to a member outside your city", None, None
        elif _same_name(match.full_name, payload["full_name"]) and match.dob == date.fromisoformat(payload["dob"]):
            outcome = "UPDATE"
        else:  # FR-78: same number, different person? Never updated automatically.
            outcome, error = "CONFLICT", "Name or date of birth differs from the member with this mobile number"
        counts[outcome] += 1
        import_rows.append(ImportRow(
            import_job_id=job.import_job_id, row_number=n, outcome=outcome, mobile_hash=mh if payload else None,
            matched_member_id=match.member_id if match else None, error_message=error,
            payload_enc=fernet.encrypt(json.dumps(payload).encode()) if payload else None))
    job.new_count, job.updated_count = counts["NEW"], counts["UPDATE"]
    job.conflict_count, job.error_count = counts["CONFLICT"], counts["ERROR"]
    db.add(job)
    db.flush()
    db.add_all(import_rows)
    db.commit()
    return _job_out(job)


def _get_job(db: Session, admin: AdminUser, job_id: UUID, lock: bool = False) -> ImportJob:
    job = db.get(ImportJob, job_id, with_for_update=lock)
    if job is None:
        raise HTTPException(404, "Import not found")
    if job.city_id != admin.city_id:
        raise audit.Denied("CROSS_CITY_DENIED", "This import belongs to another city", job_id)
    return job


def get_job(db: Session, admin: AdminUser, job_id: UUID) -> dict:
    return _job_out(_get_job(db, admin, job_id))


def list_jobs(db: Session, admin: AdminUser, limit: int, offset: int) -> dict:
    q = select(ImportJob).where(ImportJob.city_id == admin.city_id)
    total = db.scalar(select(func.count()).select_from(q.subquery()))
    jobs = db.scalars(q.order_by(ImportJob.created_at.desc()).limit(limit).offset(offset))
    return {"total": total, "items": [_job_out(j) for j in jobs]}


def _payload(row: ImportRow) -> dict | None:
    return json.loads(fernet.decrypt(row.payload_enc)) if row.payload_enc else None


def _row_out(row: ImportRow) -> dict:
    p = _payload(row)
    return {c.name: getattr(row, c.name) for c in ImportRow.__table__.columns} | {
        "full_name": p and p["full_name"], "mobile_last4": p and p["mobile"][-4:],
    }


def list_rows(db: Session, admin: AdminUser, job_id: UUID, limit: int, offset: int,
              outcome: str | None = None) -> dict:
    """The row-level report (FR-79): filter outcome=ERROR for the fixes, outcome=CONFLICT for the review list."""
    _get_job(db, admin, job_id)
    q = select(ImportRow).where(ImportRow.import_job_id == job_id)
    if outcome:
        q = q.where(ImportRow.outcome == outcome)
    total = db.scalar(select(func.count()).select_from(q.subquery()))
    rows = db.scalars(q.order_by(ImportRow.row_number).limit(limit).offset(offset))
    return {"total": total, "items": [_row_out(r) for r in rows]}


def _apply(db: Session, member: Member, p: dict, identity: bool):
    """Copy a row onto an existing member. identity=True also overwrites name and DOB (a resolved conflict)."""
    if identity:
        member.full_name, member.dob = p["full_name"], date.fromisoformat(p["dob"])
    member.gender = p["gender"]
    if p["address"]:
        member.address_enc = fernet.encrypt(p["address"].encode())
    household = db.get(Household, member.household_id) if member.household_id else None
    if household is None:
        household = Household(household_id=uuid4(), city_id=member.city_id, family_name=p["family_name"])
        db.add(household)
        db.flush()  # before the member points at it
        member.household_id, member.is_household_head = household.household_id, True
    household.family_name = p["family_name"]
    household.native_place = p["native_place"] or household.native_place


def commit(db: Session, admin: AdminUser, actor: audit.Actor, job_id: UUID) -> dict:
    """Insert NEW rows as ACTIVE members, update UPDATE rows; CONFLICT rows wait for resolve()."""
    job = _get_job(db, admin, job_id, lock=True)  # lock: a double-click can't import twice
    if job.status != "DRY_RUN_READY":
        raise HTTPException(409, f"This import is already {job.status}")
    now = datetime.now(timezone.utc)
    households, members, consents, created, applied = [], [], [], [], 0
    # ponytail: UPDATE rows load their member one query at a time; batch it if 10k-update files get slow (FR-80).
    for row in db.scalars(select(ImportRow).where(ImportRow.import_job_id == job_id,
                                                  ImportRow.outcome.in_(("NEW", "UPDATE")))).all():
        p = _payload(row)
        if row.outcome == "UPDATE":
            member = db.get(Member, row.matched_member_id) if row.matched_member_id else None
            if member is None or member.city_id != job.city_id or member.deleted_at is not None:
                row.resolution = "SKIPPED"  # gone or moved since the dry run
                continue
            _apply(db, member, p, identity=False)
            row.resolution, applied = "APPLIED", applied + 1
            continue
        h = Household(household_id=uuid4(), city_id=job.city_id, family_name=p["family_name"],
                      native_place=p["native_place"])
        # Imported by the City Admin from the community's own list, so already verified: ACTIVE.
        m = Member(member_id=uuid4(), city_id=job.city_id, household_id=h.household_id, is_household_head=True,
                   mobile_enc=fernet.encrypt(p["mobile"].encode()), mobile_hash=row.mobile_hash,
                   mobile_last4=p["mobile"][-4:], full_name=p["full_name"], gender=p["gender"],
                   dob=date.fromisoformat(p["dob"]),
                   address_enc=fernet.encrypt(p["address"].encode()) if p["address"] else None,
                   status="ACTIVE", source="IMPORT", reviewed_by_admin_id=admin.admin_id, reviewed_at=now)
        households.append(h)
        members.append(m)
        consents.append(Consent(member_id=m.member_id, purpose="MEMBERSHIP", notice_version=NOTICE_VERSION,
                                channel="IMPORT"))
        created.append((row, m.member_id))
    try:
        for batch in (households, members, consents):  # flushed in FK order; there are no relationship()s
            db.add_all(batch)
            db.flush()
        for row, member_id in created:
            row.matched_member_id, row.resolution = member_id, "APPLIED"
        job.status, job.committed_at = "COMMITTED", now
        audit.log(db, actor, "MEMBER_IMPORT", "import_job", job_id, job.city_id, None, {
            "created": len(created), "updated": applied, "conflicts": job.conflict_count, "errors": job.error_count,
        })
        db.commit()
    except IntegrityError:  # a number in the file registered through the app after the dry run
        db.rollback()
        raise HTTPException(409, "Members changed since this file was checked. Upload it again.")
    return _job_out(job)


def resolve(db: Session, admin: AdminUser, actor: audit.Actor, job_id: UUID, row_number: int, action: str) -> dict:
    """FR-78: the admin decides each conflict. APPLY overwrites the member with the row; SKIP leaves it."""
    job = _get_job(db, admin, job_id)
    if job.status not in ("DRY_RUN_READY", "COMMITTED"):
        raise HTTPException(409, f"This import is {job.status}")
    row = db.get(ImportRow, (job_id, row_number), with_for_update=True)
    if row is None:
        raise HTTPException(404, "Row not found")
    if row.outcome != "CONFLICT":
        raise HTTPException(409, "Only conflict rows need a decision")
    if row.resolution:
        raise HTTPException(409, f"This conflict is already {row.resolution}")
    if action == "SKIP":
        row.resolution = "SKIPPED"
        audit.log(db, actor, "IMPORT_CONFLICT_SKIP", "import_row", f"{job_id}:{row_number}", job.city_id)
    else:
        member = db.get(Member, row.matched_member_id, with_for_update=True) if row.matched_member_id else None
        if member is None or member.city_id != job.city_id or member.deleted_at is not None:
            raise HTTPException(409, "The matching member is no longer in your city")
        fields = ["full_name", "dob", "gender"]
        before = audit.snapshot(member, fields)
        _apply(db, member, _payload(row), identity=True)
        row.resolution = "APPLIED"
        audit.log(db, actor, "IMPORT_CONFLICT_APPLY", "member", member.member_id, job.city_id,
                  *audit.changes(before, member))
    db.commit()
    return _row_out(row)
