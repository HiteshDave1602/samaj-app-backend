"""Bulk import (SF-12) and encrypted export (SF-13) check against the dev DB: python test_import_export.py
Creates two cities, admins and members, imports files (one of 10,000 rows for FR-80), exports, then deletes it
all. audit_log rows stay: the table is append-only by design."""
import csv
import io
import secrets
import time
import zipfile
from datetime import date, datetime, timedelta, timezone
from uuid import uuid4

import msoffcrypto
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import cast, delete, func, select, String

from auth.curd import hash_token
from auth.model import AuthSession
from cityadmin.member_export.model import ExportJob
from cityadmin.member_export.schema import ALL_COLUMNS
from cityadmin.member_import.curd import HEADERS
from cityadmin.member_import.model import ImportJob
from core import private_files
from core.database import SessionLocal
from core.security import fernet, hash_password
from main import app
from member.auth.curd import mobile_hash
from member.profile.model import Consent, Household, Member
from superadmin.audit_log.model import AuditLog
from superadmin.city.model import City
from superadmin.city_admin.model import AdminUser

c = TestClient(app)
tag = secrets.token_hex(4)
now = datetime.now(timezone.utc)
mobile = lambda: f"9{secrets.randbelow(10**9):09d}"


def admin_session(db, admin: AdminUser) -> dict:
    token = secrets.token_urlsafe(32)
    db.add(AuthSession(admin_id=admin.admin_id, refresh_token_hash=hash_token(token), mfa_verified=True,
                       last_activity_at=now, expires_at=now + timedelta(hours=1)))
    return {"Authorization": f"Bearer {token}"}


def existing(db, city: City, name: str, dob: date, number: str, family: str) -> Member:
    h = Household(household_id=uuid4(), city_id=city.city_id, family_name=family)
    m = Member(member_id=uuid4(), city_id=city.city_id, household_id=h.household_id, is_household_head=True,
               mobile_hash=mobile_hash("+91" + number), mobile_enc=fernet.encrypt(("+91" + number).encode()),
               mobile_last4=number[-4:], full_name=name, gender="FEMALE", dob=dob, status="ACTIVE",
               address_enc=fernet.encrypt(b"Old address"))
    db.add(h)
    db.flush()
    db.add(m)
    db.flush()
    return m


def csv_file(rows: list[list[str]], header=HEADERS) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows([header, *rows])
    return buf.getvalue().encode()


def upload(data: bytes, name="members.csv", h=None):
    return c.post("/cityadmin/imports", headers=h or x_h, files={"file": (name, data, "application/octet-stream")})


A, B, C, N1, N2, N3 = (mobile() for _ in range(6))
with SessionLocal() as db:
    cx, cy = City(name=f"IX-{tag}", state="T"), City(name=f"IY-{tag}", state="T")
    db.add_all([cx, cy])
    db.flush()
    pw = hash_password(secrets.token_hex(8))
    sa = AdminUser(role="SUPER_ADMIN", full_name="SA", email=f"isa-{tag}@t.local", password_hash=pw)
    ca_x = AdminUser(role="CITY_ADMIN", city_id=cx.city_id, full_name="CA X", email=f"ix-{tag}@t.local", password_hash=pw)
    ca_y = AdminUser(role="CITY_ADMIN", city_id=cy.city_id, full_name="CA Y", email=f"iy-{tag}@t.local", password_hash=pw)
    db.add_all([sa, ca_x, ca_y])
    db.flush()
    sa_h, x_h, y_h = admin_session(db, sa), admin_session(db, ca_x), admin_session(db, ca_y)
    kiran = existing(db, cx, "Kiran Shah", date(1980, 1, 1), A, "Shah")
    meena = existing(db, cx, "Meena Rao", date(1990, 2, 2), B, "Rao")
    existing(db, cy, "Yash Y", date(1970, 3, 3), C, "Y family")
    db.commit()
    city_ids, admin_ids = [cx.city_id, cy.city_id], [sa.admin_id, ca_x.admin_id, ca_y.admin_id]
    kiran_id, meena_id, ca_x_uuid = kiran.member_id, meena.member_id, ca_x.admin_id
    ca_x_id = str(ca_x_uuid)

base = "/cityadmin/imports"
try:
    # template: exact headers, both formats
    r = c.get(f"{base}/template", headers=x_h, params={"file_format": "CSV"})
    assert r.status_code == 200 and r.content.decode("utf-8-sig").splitlines()[0] == ",".join(HEADERS)
    r = c.get(f"{base}/template", headers=x_h)
    assert [v.value for v in openpyxl.load_workbook(io.BytesIO(r.content)).active[1]] == HEADERS

    # FR-76: a renamed header rejects the whole file, before any write; FR-75: CSV/XLSX only
    jobs_before = c.get(base, headers=x_h).json()["total"]
    r = upload(csv_file([["N", N1, "MALE", "1985-03-03", "J", "", ""]], header=HEADERS[:-1] + ["addr"]))
    assert r.status_code == 400 and "Nothing was imported" in r.text, r.text
    assert upload(b"full_name\n", name="m.txt").status_code == 400
    assert upload(csv_file([])).status_code == 400  # no rows
    assert c.get(base, headers=x_h).json()["total"] == jobs_before

    # FR-77/78/79: dry run classifies every row; nothing changes yet
    r = upload(csv_file([
        ["Nikhil Joshi", N1, "male", "1985-03-03", "Joshi", "Pune", "4 MG Road"],  # 2 NEW (gender any case)
        ["  kiran  SHAH ", A, "FEMALE", "1980-01-01", "Shah", "Surat", "New Addr"],  # 3 UPDATE: same person
        ["Meena Iyer", B, "FEMALE", "1990-02-02", "Iyer", "", ""],  # 4 CONFLICT: same number, other name
        ["Yash Y", C, "MALE", "1970-03-03", "Y", "", ""],  # 5 ERROR: another city's member
        ["", "123", "X", "31/02/1990", "", "", ""],  # 6 ERROR: everything wrong
        ["Dup", f"+91 {N1[:5]} {N1[5:]}", "MALE", "1985-03-03", "D", "", ""],  # 7 ERROR: same number as row 2
        ["", "", "", "", "", "", ""],  # 8 blank: ignored
        ["=HYPERLINK(\"x\")", N2, "OTHER", "15/08/1975", "Desai", "Vadodara", ""],  # 9 NEW, DD/MM/YYYY
    ]))
    assert r.status_code == 201, r.text
    job = r.json()
    counts = (job["total_rows"], job["new_count"], job["updated_count"], job["conflict_count"], job["error_count"])
    assert job["status"] == "DRY_RUN_READY" and counts == (7, 2, 1, 1, 3), job
    jid = job["import_job_id"]
    rows = {r["row_number"]: r for r in c.get(f"{base}/{jid}/rows", headers=x_h).json()["items"]}
    assert [rows[n]["outcome"] for n in (2, 3, 4, 5, 6, 7, 9)] == ["NEW", "UPDATE", "CONFLICT", "ERROR", "ERROR",
                                                                   "ERROR", "NEW"], rows
    assert 8 not in rows and rows[2]["full_name"] == "Nikhil Joshi" and rows[2]["mobile_last4"] == N1[-4:]
    assert rows[3]["matched_member_id"] == str(kiran_id) and rows[4]["matched_member_id"] == str(meena_id)
    assert "outside your city" in rows[5]["error_message"] and rows[5]["matched_member_id"] is None
    assert all(s in rows[6]["error_message"] for s in ("full_name", "mobile", "gender", "dob", "family_name"))
    assert rows[7]["error_message"] == "Same mobile number as row 2"
    assert c.get(f"{base}/{jid}/rows", headers=x_h, params={"outcome": "ERROR"}).json()["total"] == 3
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).where(Member.city_id == city_ids[0])) == 2  # dry run wrote no member

    # FR-71: other city can't see or act on it
    assert c.get(f"{base}/{jid}", headers=y_h).status_code == 403
    assert c.post(f"{base}/{jid}/commit", headers=y_h).status_code == 403
    assert c.get(base, headers=y_h).json()["total"] == 0

    # commit: NEW -> ACTIVE members, UPDATE applied, CONFLICT untouched
    r = c.post(f"{base}/{jid}/commit", headers=x_h)
    assert r.status_code == 200 and r.json()["status"] == "COMMITTED" and r.json()["committed_at"], r.text
    assert c.post(f"{base}/{jid}/commit", headers=x_h).status_code == 409
    with SessionLocal() as db:
        nikhil = db.scalar(select(Member).where(Member.mobile_hash == mobile_hash("+91" + N1)))
        assert (nikhil.status, nikhil.source, nikhil.city_id, nikhil.gender) == ("ACTIVE", "IMPORT", city_ids[0], "MALE")
        assert str(nikhil.reviewed_by_admin_id) == ca_x_id and fernet.decrypt(nikhil.address_enc) == b"4 MG Road"
        hh = db.get(Household, nikhil.household_id)
        assert (hh.family_name, hh.native_place, nikhil.is_household_head) == ("Joshi", "Pune", True)
        consent = db.scalar(select(Consent).where(Consent.member_id == nikhil.member_id))
        assert (consent.purpose, consent.channel) == ("MEMBERSHIP", "IMPORT")
        k = db.get(Member, kiran_id)
        assert fernet.decrypt(k.address_enc) == b"New Addr" and k.full_name == "Kiran Shah"  # name kept as stored
        assert db.get(Household, k.household_id).native_place == "Surat"
        assert db.get(Member, meena_id).full_name == "Meena Rao"  # conflict not applied
    rows = {r["row_number"]: r for r in c.get(f"{base}/{jid}/rows", headers=x_h).json()["items"]}
    assert rows[2]["matched_member_id"] == str(nikhil.member_id) and rows[2]["resolution"] == "APPLIED"
    assert rows[4]["resolution"] is None

    # FR-78: the admin decides the conflict
    res = lambda n, action, h=x_h, j=jid: c.post(f"{base}/{j}/rows/{n}/resolve", headers=h, json={"action": action})
    assert res(4, "APPLY", y_h).status_code == 403 and res(2, "APPLY").status_code == 409
    assert res(4, "MAYBE").status_code == 422 and res(99, "SKIP").status_code == 404
    r = res(4, "APPLY")
    assert r.status_code == 200 and r.json()["resolution"] == "APPLIED", r.text
    assert res(4, "SKIP").status_code == 409
    with SessionLocal() as db:
        meena = db.get(Member, meena_id)
        assert meena.full_name == "Meena Iyer" and db.get(Household, meena.household_id).family_name == "Iyer"

    # XLSX: date cells and number cells for mobiles work; a conflict can be skipped
    wb = openpyxl.Workbook()
    wb.active.append(HEADERS)
    wb.active.append(["Kiran Kumar", int(A), "FEMALE", datetime(1980, 1, 1), "Shah", None, None])  # conflict
    wb.active.append(["Xl New", int(N3), "MALE", date(2001, 4, 4), "Xl", None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(buf.getvalue(), name="week.xlsx")
    assert r.status_code == 201 and r.json()["file_format"] == "XLSX", r.text
    assert (r.json()["conflict_count"], r.json()["new_count"]) == (1, 1), r.json()
    j2 = r.json()["import_job_id"]
    assert res(2, "SKIP", j=j2).json()["resolution"] == "SKIPPED"  # allowed during the dry run too

    # commit after a number in the file registered meanwhile: 409, nothing applied
    with SessionLocal() as db:
        existing(db, cx, "Raced In", date(2001, 4, 4), N3, "Race")
        db.commit()
    assert c.post(f"{base}/{j2}/commit", headers=x_h).status_code == 409
    assert c.get(f"{base}/{j2}", headers=x_h).json()["status"] == "DRY_RUN_READY"

    # ---- export (SF-13) ----
    ex = "/cityadmin/exports"
    pw = "Event-2026!"

    def opened(content: bytes, password: str = pw, cells: bool = False) -> list[tuple]:
        office = msoffcrypto.OfficeFile(io.BytesIO(content))
        office.load_key(password=password)
        plain = io.BytesIO()
        office.decrypt(plain)
        return [tuple(r) for r in openpyxl.load_workbook(plain).active.iter_rows(values_only=not cells)]

    assert c.post(ex, headers=x_h, json={"password": "short7!"}).status_code == 422  # FR-83: min 8
    assert c.post(ex, headers=x_h, json={"password": pw, "columns": ["biodata"]}).status_code == 422  # FR-84
    assert c.post(ex, headers=x_h, json={"password": pw, "filters": {"age_min": 50, "age_max": 40}}).status_code == 422
    r = c.post(ex, headers=x_h, json={"password": pw})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"], r.text
    assert not zipfile.is_zipfile(io.BytesIO(r.content))  # FR-83: an encrypted container, not a readable xlsx
    try:
        opened(r.content, "wrong-password")
        raise AssertionError("opened with a wrong password")
    except msoffcrypto.exceptions.InvalidKeyError:
        pass
    sheet = opened(r.content)
    assert list(sheet[0]) == ALL_COLUMNS and int(r.headers["x-row-count"]) == len(sheet) - 1, sheet
    names = {row[0] for row in sheet[1:]}  # own city only, ACTIVE by default
    assert names == {"Kiran Shah", "Meena Iyer", "Nikhil Joshi", "=HYPERLINK(\"x\")", "Raced In"}, names
    formula_like = [row[0] for row in opened(r.content, cells=True) if str(row[0].value).startswith("=")]
    assert [cell.data_type for cell in formula_like] == ["s"]  # stored as text, never run as a formula

    # filters and columns
    r = c.post(ex, headers=x_h, json={"password": pw, "columns": ["full_name", "mobile", "age"],
                                      "filters": {"gender": "FEMALE", "age_min": 30, "age_max": 50}})
    sheet = opened(r.content)
    assert sheet == [("full_name", "mobile", "age"), ("Kiran Shah", "+91" + A, date.today().year - 1980),
                     ("Meena Iyer", "+91" + B, date.today().year - 1990 - (date.today() < date(date.today().year, 2, 2)))], sheet
    sheet = opened(c.post(ex, headers=x_h, json={"password": pw, "columns": ["full_name"],
                                                 "filters": {"area": "mg road"}}).content)
    assert sheet == [("full_name",), ("Nikhil Joshi",)], sheet
    assert opened(c.post(ex, headers=y_h, json={"password": pw, "columns": ["full_name"]}).content) == [
        ("full_name",), ("Yash Y",)]
    r = c.post(ex, headers=x_h, json={"password": pw, "columns": ["full_name"], "filters": {"status": ["PENDING"]}})
    assert r.headers["x-row-count"] == "0"

    # history, and FR-85: the password is in no stored row
    hist = c.get(ex, headers=x_h).json()
    assert hist["total"] == 4 and hist["items"][0]["filters"]["status"] == ["PENDING"], hist
    assert hist["items"][-1]["row_count"] == 5 and hist["items"][-1]["columns_exported"] == ALL_COLUMNS
    assert c.get(ex, headers=y_h).json()["total"] == 1
    with SessionLocal() as db:
        leaked = db.scalar(select(func.count()).where(AuditLog.actor_id == ca_x_uuid,
                                                      cast(AuditLog.after_state, String).contains(pw)))
        leaked += db.scalar(select(func.count()).where(cast(ExportJob.filters, String).contains(pw)))
        assert leaked == 0

    # FR-80: 10,000 rows validated, deduplicated and committed within 2 minutes
    prefix = f"7{secrets.randbelow(10**4):04d}"
    big = csv_file([[f"Bulk {i}", f"{prefix}{i:05d}", "MALE", "1990-01-01", "Bulk", "", ""] for i in range(10_000)])
    started = time.monotonic()
    r = upload(big)
    assert r.status_code == 201 and r.json()["new_count"] == 10_000, r.text[:300]
    assert c.post(f"{base}/{r.json()['import_job_id']}/commit", headers=x_h).status_code == 200
    took = time.monotonic() - started
    assert took < 120, took

    actions = sorted(e["action"] for e in c.get("/superadmin/audit-logs", headers=sa_h,
                                                params={"actor_id": ca_x_id, "limit": 200}).json()["items"])
    assert actions == ["IMPORT_CONFLICT_APPLY", "IMPORT_CONFLICT_SKIP"] + ["MEMBER_EXPORT"] * 4 + ["MEMBER_IMPORT"] * 2, actions
    print(f"ok (10,000-row import + commit: {took:.1f} s)")
finally:
    with SessionLocal() as db:  # FK order; import rows go via ON DELETE CASCADE
        file_keys = db.scalars(select(ImportJob.file_key).where(ImportJob.city_id.in_(city_ids))).all()
        member_ids = select(Member.member_id).where(Member.city_id.in_(city_ids))
        db.execute(delete(ImportJob).where(ImportJob.city_id.in_(city_ids)))
        db.execute(delete(ExportJob).where(ExportJob.city_id.in_(city_ids)))
        db.execute(delete(Consent).where(Consent.member_id.in_(member_ids)))
        db.execute(delete(Member).where(Member.city_id.in_(city_ids)))
        db.execute(delete(Household).where(Household.city_id.in_(city_ids)))
        db.execute(delete(AdminUser).where(AdminUser.admin_id.in_(admin_ids)))
        db.execute(delete(City).where(City.city_id.in_(city_ids)))
        db.commit()
    for key in file_keys:
        private_files.delete(key)
